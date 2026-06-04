"""
Flask web application for CCS Annotation Tool.
"""
import os, sys, uuid, threading, csv, io, tempfile, json, time, shutil
from flask import Flask, request, jsonify, render_template, send_file, make_response
from werkzeug.utils import secure_filename

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from processor import (
    process_single_mode, load_chemclass, load_kegg, open_db,
    simplify_name, determine_best_adduct, norm_adduct, parse_adduct_mass_diff,
    ADDUCT_MASSES, rank_compound_rows, hmdb_dedup, extract_raw_matrix
)
from metabolomics import run_analysis

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 2 * 1024 * 1024 * 1024  # 2GB max upload

TASKS = {}
TEMP_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'temp')
os.makedirs(TEMP_DIR, exist_ok=True)

# ============================================================
# Helpers
# ============================================================
def parse_hmdb_ids(text_or_rows):
    """Parse HMDB IDs from paste text or list of rows. Returns list of IDs."""
    if isinstance(text_or_rows, str):
        ids = []
        for line in text_or_rows.strip().split('\n'):
            for part in line.replace(',', ';').split(';'):
                part = part.strip()
                if part: ids.append(part)
        return ids
    # list of rows from file import
    ids = []
    # Try to find HMDB ID column
    if not text_or_rows: return ids
    header = text_or_rows[0]
    hmdb_col = None
    for i, h in enumerate(header):
        hl = h.lower().strip()
        if 'hmdb' in hl or 'compound id' in hl or ('id' in hl and 'hmdb' not in hl):
            hmdb_col = i; break
    if hmdb_col is None:
        hmdb_col = 0  # Default to first column
    for row in text_or_rows[1:]:
        if row and len(row) > hmdb_col:
            ids.append(row[hmdb_col].strip())
    return ids

# ============================================================
# Routes
# ============================================================
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/origins_tissues')
def api_origins_tissues():
    """Return available Origin and Tissue values for checkboxes."""
    conn = open_db()
    c = conn.cursor()
    origins = set()
    tissues = set()
    c.execute('SELECT origin, biospecimens, tissues FROM hmdb_source')
    for o, b, t in c.fetchall():
        if o:
            for v in (o or '').replace(';', ',').split(','):
                v = v.strip()
                if v: origins.add(v)
        if b:
            for v in (b or '').replace(';', ',').split(','):
                v = v.strip()
                if v: tissues.add(v)
        if t:
            for v in (t or '').replace(';', ',').split(','):
                v = v.strip()
                if v: tissues.add(v)
    conn.close()
    return jsonify({'origins': sorted(origins), 'tissues': sorted(tissues)})

# --- Tab 1: Full pipeline ---
@app.route('/api/run', methods=['POST'])
def api_run():
    task_id = uuid.uuid4().hex
    task_dir = os.path.join(TEMP_DIR, task_id)
    os.makedirs(task_dir, exist_ok=True)

    # Save uploaded files
    files = {}
    for key in ['pos_com', 'pos_iden', 'neg_com', 'neg_iden']:
        if key in request.files:
            f = request.files[key]
            if f.filename:
                fname = secure_filename(f.filename)
                fpath = os.path.join(task_dir, fname)
                f.save(fpath)
                files[key] = fpath

    if 'pos_com' not in files or 'pos_iden' not in files:
        return jsonify({'error': 'POS COM and IDEN are required'}), 400

    # Parse filter params
    def parse_float(val):
        v = val.strip() if val else ''
        return float(v) if v else None
    def parse_int(val):
        v = val.strip() if val else ''
        return int(v) if v else None

    params = {
        'score_cutoff': parse_int(request.form.get('score_cutoff', '36')),
        'frag_cutoff': parse_float(request.form.get('frag_cutoff', '1')),
        'dccs_cutoff': parse_float(request.form.get('dccs_cutoff', '')),
        'allowed_origins': request.form.get('allowed_origins', '').split(',') if request.form.get('allowed_origins') else [],
        'priority_tissues': set(request.form.get('priority_tissues', '').split(',')) if request.form.get('priority_tissues') else set(),
    }

    log_lines = []
    def log(msg):
        log_lines.append(msg)

    def run_pipeline():
        try:
            log("=== Starting CCS Annotation Pipeline ===")
            t0 = time.time()
            all_filtered = []
            pos_h1 = pos_nh = neg_h1 = neg_nh = None
            pos_rows_list = neg_rows_list = None

            # POS
            log("=== POS Mode ===")
            pos_result = process_single_mode(
                files['pos_com'], files['pos_iden'],
                params['score_cutoff'], params['frag_cutoff'], params['dccs_cutoff'],
                params['allowed_origins'], params['priority_tissues'], log
            )
            if pos_result:
                pos_rows, pos_h1, pos_nh, pos_summary = pos_result
                pos_rows_list = pos_rows
                all_filtered.extend(pos_rows)
                # Write POS output
                pos_out = os.path.join(task_dir, 'output_POS.csv')
                with open(pos_out, 'w', newline='', encoding='utf-8-sig') as f:
                    w = csv.writer(f)
                    if pos_h1 is not None: w.writerow(pos_h1)
                    w.writerow(pos_nh); w.writerows(pos_rows)
                log(f"  POS saved: output_POS.csv ({len(pos_rows)} compounds)")

            # NEG
            has_neg = 'neg_com' in files and 'neg_iden' in files
            if has_neg:
                log("=== NEG Mode ===")
                neg_result = process_single_mode(
                    files['neg_com'], files['neg_iden'],
                    params['score_cutoff'], params['frag_cutoff'], params['dccs_cutoff'],
                    params['allowed_origins'], params['priority_tissues'], log
                )
                if neg_result:
                    neg_rows, neg_h1, neg_nh, neg_summary = neg_result
                    neg_rows_list = neg_rows
                    all_filtered.extend(neg_rows)
                    neg_out = os.path.join(task_dir, 'output_NEG.csv')
                    with open(neg_out, 'w', newline='', encoding='utf-8-sig') as f:
                        w = csv.writer(f)
                        if neg_h1 is not None: w.writerow(neg_h1)
                        w.writerow(neg_nh); w.writerows(neg_rows)
                    log(f"  NEG saved: output_NEG.csv ({len(neg_rows)} compounds)")

            # Merge
            if has_neg and pos_rows_list and neg_rows_list:
                log("=== Merging POS+NEG by Compound ID ===")
                id_col, score_col = 1, 5
                groups = {}
                no_id_rows = []
                for row in all_filtered:
                    cid = (row[id_col] or '').strip().upper() if id_col < len(row) else ''
                    if cid:
                        if cid not in groups: groups[cid] = []
                        groups[cid].append(row)
                    else:
                        no_id_rows.append(row)

                merged = list(no_id_rows)
                for cid, rows in groups.items():
                    best = max(rows, key=lambda r: float(r[score_col]) if score_col < len(r) and r[score_col] else 0)
                    merged.append(best)

                log(f"  POS: {pos_summary['filtered']}  NEG: {neg_summary['filtered']}  Merged: {len(merged)}")

                # Normalize columns
                ncols = len(pos_nh)
                for row in merged:
                    while len(row) < ncols: row.append('')
                    if len(row) > ncols: del row[ncols:]

                merge_out = os.path.join(task_dir, 'output_merged.csv')
                with open(merge_out, 'w', newline='', encoding='utf-8-sig') as f:
                    w = csv.writer(f)
                    if pos_h1 is not None: w.writerow(pos_h1)
                    w.writerow(pos_nh); w.writerows(merged)
                log(f"  Merged saved: output_merged.csv ({len(merged)} compounds)")

                # Raw matrix
                matrix_header, matrix_rows = extract_raw_matrix(merged, pos_nh)
                matrix_out = os.path.join(task_dir, 'output_merged_raw_matrix.csv')
                with open(matrix_out, 'w', newline='', encoding='utf-8-sig') as f:
                    w = csv.writer(f)
                    w.writerow(matrix_header); w.writerows(matrix_rows)
                log(f"  Raw matrix saved: output_merged_raw_matrix.csv ({len(matrix_rows)} rows x {len(matrix_header)-1} cols)")

            # Raw matrix for POS-only
            elif pos_rows_list:
                matrix_header, matrix_rows = extract_raw_matrix(pos_rows_list, pos_nh)
                matrix_out = os.path.join(task_dir, 'output_POS_raw_matrix.csv')
                with open(matrix_out, 'w', newline='', encoding='utf-8-sig') as f:
                    w = csv.writer(f)
                    w.writerow(matrix_header); w.writerows(matrix_rows)

            elapsed = time.time() - t0
            log(f"\n=== Done ({elapsed:.0f}s) ===")

        except Exception as e:
            log(f"\nERROR: {e}")
            import traceback
            log(traceback.format_exc())

    thread = threading.Thread(target=run_pipeline, daemon=True)
    TASKS[task_id] = {'thread': thread, 'log': log_lines, 'dir': task_dir, 'done': False}
    thread.start()
    return jsonify({'task_id': task_id})

@app.route('/api/status/<task_id>')
def api_status(task_id):
    task = TASKS.get(task_id)
    if not task:
        return jsonify({'error': 'Task not found'}), 404
    files = []
    if task['thread'].is_alive() is False and not task.get('done'):
        task['done'] = True
        for f in os.listdir(task['dir']):
            if f.endswith('.csv'):
                files.append(f)
    return jsonify({
        'running': task['thread'].is_alive(),
        'log': '\n'.join(task.get('log', [])),
        'files': files
    })

@app.route('/api/metabolomics/run', methods=['POST'])
def api_metabolomics_run():
    task_id = uuid.uuid4().hex
    task_dir = os.path.join(TEMP_DIR, task_id)
    os.makedirs(task_dir, exist_ok=True)

    if 'matrix' not in request.files or not request.files['matrix'].filename:
        return jsonify({'error': 'A metabolomics matrix CSV is required'}), 400

    matrix_file = request.files['matrix']
    matrix_name = secure_filename(matrix_file.filename) or 'matrix.csv'
    matrix_path = os.path.join(task_dir, matrix_name)
    matrix_file.save(matrix_path)

    group_a = request.form.get('group_a') or None
    group_b = request.form.get('group_b') or None
    log_lines = []

    def log(msg):
        log_lines.append(msg)

    def run_job():
        try:
            log("=== Starting metabolomics analysis ===")
            t0 = time.time()
            result = run_analysis(matrix_path, task_dir, group_a, group_b)
            TASKS[task_id]['result'] = result
            log(f"  Matrix: {result['summary']['feature_count']} features x {result['summary']['sample_count']} samples")
            log(f"  Groups: " + ', '.join(f"{g['group']}({g['count']})" for g in result['summary']['groups']))
            log(f"  Comparison: {result['comparison']['group_a']} vs {result['comparison']['group_b']}")
            log(f"  Significant features: {result['significant_count']} (FDR <= 0.05 and |log2FC| >= 1)")
            for warning in result['summary']['warnings']:
                log(f"  Warning: {warning}")
            log(f"=== Done ({time.time() - t0:.1f}s) ===")
        except Exception as e:
            TASKS[task_id]['error'] = str(e)
            log(f"ERROR: {e}")
            import traceback
            log(traceback.format_exc())

    thread = threading.Thread(target=run_job, daemon=True)
    TASKS[task_id] = {'thread': thread, 'log': log_lines, 'dir': task_dir, 'done': False}
    thread.start()
    return jsonify({'task_id': task_id})

@app.route('/api/metabolomics/status/<task_id>')
def api_metabolomics_status(task_id):
    task = TASKS.get(task_id)
    if not task:
        return jsonify({'error': 'Task not found'}), 404
    files = []
    if task['thread'].is_alive() is False:
        for f in os.listdir(task['dir']):
            if f.endswith('.csv'):
                files.append(f)
    return jsonify({
        'running': task['thread'].is_alive(),
        'log': '\n'.join(task.get('log', [])),
        'files': files,
        'result': task.get('result'),
        'error': task.get('error'),
    })

@app.route('/api/download/<task_id>/<filename>')
def api_download(task_id, filename):
    task = TASKS.get(task_id)
    if not task: return jsonify({'error': 'Task not found'}), 404
    filepath = os.path.join(task['dir'], filename)
    if not os.path.exists(filepath): return jsonify({'error': 'File not found'}), 404
    return send_file(filepath, as_attachment=True, download_name=filename)

# --- Tab 2: HMDB Source Query ---
@app.route('/api/hmdb_query', methods=['POST'])
def api_hmdb_query():
    ids = []
    if 'file' in request.files and request.files['file'].filename:
        f = request.files['file']
        if f.filename.endswith('.csv'):
            content = f.read().decode('utf-8-sig', errors='replace')
            rows = list(csv.reader(io.StringIO(content)))
            ids = parse_hmdb_ids(rows)
        elif f.filename.endswith(('.xlsx', '.xls')):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
                ws = wb.active
                rows = [[str(c.value or '') for c in row] for row in ws.iter_rows()]
                ids = parse_hmdb_ids(rows)
            except: pass
    else:
        text = request.form.get('text', '')
        ids = parse_hmdb_ids(text)

    if not ids:
        return jsonify({'error': 'No HMDB IDs found'}), 400

    conn = open_db(); c = conn.cursor()
    results = []
    for hid in ids:
        hid_upper = hid.upper()
        c.execute('SELECT origin, biospecimens, tissues, cellular FROM hmdb_source WHERE hmdb_id=?', (hid_upper,))
        row = c.fetchone()
        if row:
            results.append([hid, row[0] or '', row[1] or '', row[2] or '', row[3] or ''])
        else:
            results.append([hid, '', '', '', ''])
    conn.close()

    return jsonify({'columns': ['HMDB_ID', 'Origin', 'Biospecimens', 'Tissues', 'Cellular'], 'rows': results})

# --- Tab 3: CCS Value Query ---
@app.route('/api/ccs_query', methods=['POST'])
def api_ccs_query():
    pairs = []
    if 'file' in request.files and request.files['file'].filename:
        f = request.files['file']
        if f.filename.endswith('.csv'):
            content = f.read().decode('utf-8-sig', errors='replace')
            rows = list(csv.reader(io.StringIO(content)))
            for row in rows:
                if len(row) >= 2:
                    pairs.append((row[0].strip(), row[1].strip()))
        elif f.filename.endswith(('.xlsx', '.xls')):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
                ws = wb.active
                for rrow in ws.iter_rows(min_row=2):
                    vals = [str(c.value or '').strip() for c in rrow]
                    if len(vals) >= 2 and vals[0] and vals[1]:
                        pairs.append((vals[0], vals[1]))
            except: pass
    else:
        text = request.form.get('text', '')
        for line in text.strip().split('\n'):
            parts = [p.strip() for p in line.replace('\t', ',').split(',')]
            if len(parts) >= 2:
                pairs.append((parts[0], parts[1]))

    if not pairs:
        return jsonify({'error': 'No HMDB+Adduct pairs found'}), 400

    conn = open_db(); c = conn.cursor()
    results = []
    for hid, adduct in pairs:
        na = norm_adduct(adduct)
        c.execute("SELECT ccs, ccs_type, source_db, name FROM ccs WHERE hmdb_id=? AND adduct LIKE ? LIMIT 5",
                  (hid.upper(), f'%{na}%'))
        rows = c.fetchall()
        if rows:
            best = min(rows, key=lambda r: r[0] or 999)
            results.append([hid, adduct, str(best[0]) if best[0] else '', best[1] or '', best[2] or '', best[3] or ''])
        else:
            c.execute("SELECT ccs, ccs_type, source_db, name FROM ccs WHERE hmdb_id=? LIMIT 3", (hid.upper(),))
            rows = c.fetchall()
            if rows:
                best = max(rows, key=lambda r: 1 if r[1] == 'Experimental' else 0)
                results.append([hid, adduct, str(best[0]) if best[0] else '', best[1] or '', best[2] or '', best[3] or ''])
            else:
                results.append([hid, adduct, '', '', '', ''])
    conn.close()

    return jsonify({'columns': ['HMDB_ID', 'Adduct', 'CCS', 'CCS_Type', 'Source_DB', 'Name'], 'rows': results})

# --- Tab 4: Origin Query ---
@app.route('/api/origin_query', methods=['POST'])
def api_origin_query():
    ids = []
    if 'file' in request.files and request.files['file'].filename:
        f = request.files['file']
        if f.filename.endswith('.csv'):
            content = f.read().decode('utf-8-sig', errors='replace')
            rows = list(csv.reader(io.StringIO(content)))
            ids = parse_hmdb_ids(rows)
        elif f.filename.endswith(('.xlsx', '.xls')):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
                ws = wb.active
                rows = [[str(c.value or '') for c in row] for row in ws.iter_rows()]
                ids = parse_hmdb_ids(rows)
            except: pass
    else:
        ids = parse_hmdb_ids(request.form.get('text', ''))

    if not ids:
        return jsonify({'error': 'No HMDB IDs found'}), 400

    conn = open_db(); c = conn.cursor()
    results = []
    for hid in ids:
        c.execute('SELECT origin FROM hmdb_source WHERE hmdb_id=?', (hid.upper(),))
        row = c.fetchone()
        results.append([hid, row[0] if row else ''])
    conn.close()

    return jsonify({'columns': ['HMDB_ID', 'Origin'], 'rows': results})

# --- Tab 5: Chemical Class Query ---
@app.route('/api/chemclass_query', methods=['POST'])
def api_chemclass_query():
    ids = []
    if 'file' in request.files and request.files['file'].filename:
        f = request.files['file']
        if f.filename.endswith('.csv'):
            content = f.read().decode('utf-8-sig', errors='replace')
            rows = list(csv.reader(io.StringIO(content)))
            ids = parse_hmdb_ids(rows)
        elif f.filename.endswith(('.xlsx', '.xls')):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
                ws = wb.active
                rows = [[str(c.value or '') for c in row] for row in ws.iter_rows()]
                ids = parse_hmdb_ids(rows)
            except: pass
    else:
        ids = parse_hmdb_ids(request.form.get('text', ''))

    if not ids:
        return jsonify({'error': 'No HMDB IDs found'}), 400

    chemclass = load_chemclass()
    results = [[hid, chemclass.get(hid.upper(), '')] for hid in ids]

    return jsonify({'columns': ['HMDB_ID', 'Chemical_Class'], 'rows': results})

# --- Tab 6: KEGG ID Query ---
@app.route('/api/kegg_query', methods=['POST'])
def api_kegg_query():
    ids = []
    if 'file' in request.files and request.files['file'].filename:
        f = request.files['file']
        if f.filename.endswith('.csv'):
            content = f.read().decode('utf-8-sig', errors='replace')
            rows = list(csv.reader(io.StringIO(content)))
            ids = parse_hmdb_ids(rows)
        elif f.filename.endswith(('.xlsx', '.xls')):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
                ws = wb.active
                rows = [[str(c.value or '') for c in row] for row in ws.iter_rows()]
                ids = parse_hmdb_ids(rows)
            except: pass
    else:
        ids = parse_hmdb_ids(request.form.get('text', ''))

    if not ids:
        return jsonify({'error': 'No HMDB IDs found'}), 400

    kegg = load_kegg()
    results = [[hid, kegg.get(hid.upper(), '')] for hid in ids]

    return jsonify({'columns': ['HMDB_ID', 'KEGG_ID'], 'rows': results})

# --- Export CSV helper ---
@app.route('/api/export_csv', methods=['POST'])
def api_export_csv():
    data = request.json
    if not data: return jsonify({'error': 'No data'}), 400
    output = io.StringIO()
    w = csv.writer(output)
    w.writerow(data.get('columns', []))
    w.writerows(data.get('rows', []))
    resp = make_response(output.getvalue())
    resp.headers['Content-Type'] = 'text/csv; charset=utf-8-sig'
    resp.headers['Content-Disposition'] = 'attachment; filename=export.csv'
    return resp

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True, threaded=True)
